# service/diagnostic/diag_fill_volodarskogo.py
"""Диагностика заливок в многостраничном Excel Володарского."""
import sys
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parents[2]
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from openpyxl import load_workbook


def main():
    if len(sys.argv) > 1:
        fp = Path(sys.argv[1])
    else:
        fp = BASE_DIR / "data" / "input" / "Володарского 18 октябрь" / "График дежурств 2026.xlsx"

    print(f"Читаю: {fp}")
    if not fp.exists():
        print("❌ Файл не найден")
        return

    wb = load_workbook(fp, data_only=True)
    print(f"Листы: {wb.sheetnames}")
    print()

    # Возьмём конкретный лист (октябрь) — или первый, если его нет
    target_sheet = "октябрь" if "октябрь" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[target_sheet]
    print(f"=== Анализ листа: {target_sheet} ===")
    print(f"Строк: {ws.max_row}, колонок: {ws.max_column}")
    print()

    color_stats: dict[str, int] = {}
    colored_cells: list[tuple[str, str, str]] = []

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            fill = cell.fill
            if fill is None or fill.fill_type is None:
                continue
            rgb = getattr(fill.start_color, 'rgb', None)
            if not rgb:
                continue
            if rgb in ('00000000', 'FFFFFFFF'):
                continue
            color_stats[rgb] = color_stats.get(rgb, 0) + 1
            if len(colored_cells) < 50:
                colored_cells.append((cell.coordinate, str(cell.value)[:40], rgb))

    print(f"Цветных ячеек: {sum(color_stats.values())}")
    print(f"Уникальных цветов: {len(color_stats)}")
    for rgb, count in sorted(color_stats.items(), key=lambda x: -x[1]):
        print(f"  {rgb}: {count}")

    print()
    print("Первые 50 цветных ячеек (координата, значение, RGB):")
    for coord, val, rgb in colored_cells:
        print(f"  {coord}: {val!r}  [{rgb}]")


if __name__ == "__main__":
    main()