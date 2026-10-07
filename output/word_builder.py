# output/word_builder.py
"""Сохранение сводного графика в формат Word (.docx)."""
import calendar
import datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from core.config import MONTHS_RU, MONTHS_RU_UPPER, DEPT_ABBR
from core.utils import normalize_doctor_name_for_output, round_time_up_for_output

def get_dept_abbr_from_name(dept_name: str) -> str:
    """Возвращает аббревиатуру отделения."""
    if dept_name in DEPT_ABBR:
        return DEPT_ABBR[dept_name]
    for key, abbr in DEPT_ABBR.items():
        if key.lower() in dept_name.lower() or dept_name.lower() in key.lower():
            return abbr
    words = dept_name.split()
    if len(words) >= 2:
        return ''.join(w[0].upper() for w in words[:2])
    return dept_name[:4].upper()

def save_to_word(all_data, doctors_by_dept, sorted_depts, output_file,
                 rukovoditel_text, ploshadka_text, month_num, year):
    """Сохраняет сводный график в Word с тремя колонками и шапкой."""
    doc = Document()
    
    section = doc.sections[0]
    section.top_margin = Cm(0.8)
    section.bottom_margin = Cm(0.8)
    section.left_margin = Cm(0.8)
    section.right_margin = Cm(0.8)

    # ШАПКА
    ruk_parts = rukovoditel_text.split('|')
    ruk_dolzhnost = ruk_parts[0].strip() if len(ruk_parts) > 0 else ''
    ruk_fio = ruk_parts[1].strip() if len(ruk_parts) > 1 else ''

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run("УТВЕРЖДАЮ")
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    run.font.bold = True

    if ruk_dolzhnost:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(ruk_dolzhnost)
        run.font.name = 'Calibri'
        run.font.size = Pt(10)

    if ruk_fio:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(f"____________ {ruk_fio}")
        run.font.name = 'Calibri'
        run.font.size = Pt(10)

    today = datetime.datetime.now()
    month_ru = MONTHS_RU.get(today.month, 'августа')
    date_str = f"{today.day} {month_ru} {today.year} г."
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(date_str)
    run.font.name = 'Calibri'
    run.font.size = Pt(10)

    doc.add_paragraph()

    # ЗАГОЛОВОК
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run("ГРАФИК ДЕЖУРСТВА ВРАЧЕЙ СТАЦИОНАРА ГОБУЗ МОКМЦ")
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    run.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(ploshadka_text.upper())
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    run.font.bold = True

    month_upper = MONTHS_RU_UPPER.get(month_num, 'АВГУСТ')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(f"НА {month_upper} {year} Г.")
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    run.font.bold = True

    doc.add_paragraph()

    # ДНИ НЕДЕЛИ
    _, last_day = calendar.monthrange(year, month_num)
    weekday_names = ['Понедельник', 'Вторник', 'Среда', 'Четверг', 'Пятница', 'Суббота', 'Воскресенье']
    weekdays = [weekday_names[datetime.date(year, month_num, day).weekday()] for day in range(1, last_day + 1)]

    # СБОРКА БЛОКОВ (ИСПРАВЛЕННЫЙ ФОРМАТ ОАР)
    all_blocks = []
    for day in sorted(all_data.keys()):
        if 1 <= day <= last_day:
            day_name = weekdays[day-1]
            block_lines = [f"{day:02d}.{month_num:02d} {day_name}"]
            for dept in sorted_depts:
                dept_doctors = sorted(doctors_by_dept.get(dept, []), key=lambda x: x[0])
                abbr = get_dept_abbr_from_name(dept)
                is_travma_zav = (dept == 'Травматология (зав.)')

                for doctor, dept_key in dept_doctors:
                    if (doctor, dept_key) in all_data[day]:
                        value = all_data[day][(doctor, dept_key)]

                        # Нормализация ФИО для травматологии заведующего
                        out_doctor = normalize_doctor_name_for_output(doctor) if is_travma_zav else doctor

                        # Администрация — без времени
                        if dept == 'АДМ' or dept_key == 'АДМ':
                            block_lines.append(f"{out_doctor} ({abbr})")
                            continue

                        # Парсим значение: если начинается с А/Р/Э — буква после скобки
                        letter = ''
                        time_part = value
                        if value and value[0] in 'АРЭ':
                            letter = value[0]
                            time_part = value[1:]

                        # Округление времени для травматологии заведующего
                        if is_travma_zav:
                            time_part = round_time_up_for_output(time_part)

                        if letter:
                            block_lines.append(f"{out_doctor} ({abbr}){letter} {time_part}")
                        else:
                            block_lines.append(f"{out_doctor} ({abbr}) {time_part}")
            all_blocks.append(block_lines)

    # РАСПРЕДЕЛЕНИЕ ПО СТРАНИЦАМ
    def calc_rows_for_col(blocks):
        total = 0
        for block_idx, block in enumerate(blocks):
            total += len(block)
            if block_idx < len(blocks) - 1:
                total += 1
        return total

    FIRST_PAGE_MAX_ROWS = 48
    OTHER_PAGE_MAX_ROWS = 59

    def distribute_blocks_with_limit(blocks, first_max, other_max):
        pages = []
        current_page = [[], [], []]
        col_heights = [0, 0, 0]
        current_col = 0
        is_first_page = True

        for block in blocks:
            block_rows = len(block) + 1
            max_rows = first_max if is_first_page else other_max

            if col_heights[current_col] + block_rows <= max_rows:
                current_page[current_col].append(block)
                col_heights[current_col] += block_rows
            else:
                col_found = False
                for next_col in range(current_col + 1, 3):
                    if col_heights[next_col] + block_rows <= max_rows:
                        current_col = next_col
                        current_page[current_col].append(block)
                        col_heights[current_col] += block_rows
                        col_found = True
                        break
                
                if not col_found:
                    if any(len(col) > 0 for col in current_page):
                        pages.append(current_page)
                    current_page = [[], [], []]
                    col_heights = [0, 0, 0]
                    current_col = 0
                    is_first_page = False
                    current_page[current_col].append(block)
                    col_heights[current_col] += block_rows

        if any(len(col) > 0 for col in current_page):
            pages.append(current_page)
        return pages

    pages = distribute_blocks_with_limit(all_blocks, FIRST_PAGE_MAX_ROWS, OTHER_PAGE_MAX_ROWS)
    if not pages:
        pages = [[[], [], []]]

    def fill_column(table, col_idx, blocks, start_row, max_rows):
        row = start_row
        for block_idx, block in enumerate(blocks):
            if row >= max_rows:
                break
            cell = table.cell(row, col_idx)
            cell.paragraphs[0].clear()
            p = cell.paragraphs[0]
            run = p.add_run(block[0])
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            run.font.bold = True
            row += 1

            for line in block[1:]:
                if row >= max_rows:
                    break
                cell = table.cell(row, col_idx)
                cell.paragraphs[0].clear()
                p = cell.paragraphs[0]
                run = p.add_run(line)
                run.font.name = 'Calibri'
                run.font.size = Pt(10)
                
                # <-- ДОБАВЛЕНО: Проверка на администрацию для жирного шрифта
                # Жирный шрифт для администрации
                if "(АДМ)" in line:
                    run.font.bold = True
                else:
                    run.font.bold = False
                    
                row += 1

            is_last_block = (block_idx == len(blocks) - 1)
            if not is_last_block and row < max_rows:
                cell = table.cell(row, col_idx)
                cell.paragraphs[0].clear()
                row += 1
        return row

    first_page = True
    for page_idx, page in enumerate(pages):
        if not any(len(col) > 0 for col in page):
            continue
        
        if not first_page:
            doc.add_page_break()
        first_page = False

        max_rows = max(
            calc_rows_for_col(page[0]),
            calc_rows_for_col(page[1]),
            calc_rows_for_col(page[2])
        )
        table_height = max(max_rows, 3)
        table = doc.add_table(rows=table_height, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for row in table.rows:
            for cell in row.cells:
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                tcBorders = OxmlElement('w:tcBorders')
                for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                    border = OxmlElement(f'w:{border_name}')
                    border.set(qn('w:val'), 'none')
                    border.set(qn('w:sz'), '0')
                    tcBorders.append(border)
                tcPr.append(tcBorders)
                cell.paragraphs[0].paragraph_format.space_before = Pt(0)
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                cell.paragraphs[0].paragraph_format.line_spacing = 1.0

        fill_column(table, 0, page[0], 0, table_height)
        fill_column(table, 1, page[1], 0, table_height)
        fill_column(table, 2, page[2], 0, table_height)

        for row in table.rows:
            for cell in row.cells:
                cell.width = Inches(2.5)

    try:
        doc.save(output_file)
        print(f"💾 Word сохранён: {output_file}")
    except PermissionError:
        raise PermissionError(f"Файл '{Path(output_file).name}' открыт в Word! Закройте его и повторите попытку.")